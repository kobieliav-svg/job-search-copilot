#!/usr/bin/env python3
"""Fill the user's real cover letter template with tailored content for one job posting.

If templates/cover_letter_master.docx exists, this edits it in place (same philosophy
as generate_resume.py: preserve the user's real formatting). If the user never gave us
a cover letter sample, it builds a plain, simply-formatted one from scratch instead --
cover letters carry a lot less formatting risk than resumes.

Content JSON shape:
{
  "placeholders": {
    "{{COMPANY}}": "Acme Inc.",
    "{{ROLE}}": "Senior Data Analyst",
    "{{DATE}}": "July 24, 2026"
  },
  "paragraph_replacements": {
    "Opening": "Tailored opening paragraph...",
    "Body": "Tailored body paragraph...",
    "Closing": "Tailored closing paragraph..."
  },
  "fresh_build": {
    "salutation": "Dear Hiring Manager,",
    "paragraphs": ["Opening paragraph...", "Body paragraph...", "Closing paragraph..."],
    "sign_off": "Sincerely,",
    "name": "Your Name"
  }
}

"fresh_build" is only used when there's no real template docx to edit in place.

Usage:
    python scripts/generate_cover_letter.py --template templates/cover_letter_master.docx \\
        --content output/tmp_acme_cl.json --output "output/Acme_DataAnalyst_CoverLetter.docx"
"""
import argparse
import json
import sys
from pathlib import Path

from docx import Document

import docx_editor


def build_fresh(fresh: dict) -> Document:
    doc = Document()
    doc.add_paragraph(fresh.get("salutation", "Dear Hiring Manager,"))
    for para in fresh.get("paragraphs", []):
        doc.add_paragraph(para)
    doc.add_paragraph(fresh.get("sign_off", "Sincerely,"))
    doc.add_paragraph(fresh.get("name", ""))
    return doc


def apply_content(doc: Document, content: dict) -> list[str]:
    warnings = []
    for heading, text in content.get("paragraph_replacements", {}).items():
        if not docx_editor.replace_paragraph_text(doc, heading, text):
            warnings.append(f"Could not find heading '{heading}' for a paragraph replacement.")
    for placeholder, value in content.get("placeholders", {}).items():
        n = docx_editor.replace_placeholder(doc, placeholder, value)
        if n == 0:
            warnings.append(f"Placeholder '{placeholder}' not found anywhere in the document.")
    return warnings


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("--template", required=True)
    ap.add_argument("--content", required=True)
    ap.add_argument("--output", required=True)
    args = ap.parse_args()

    template_path = Path(args.template)
    content_path = Path(args.content)
    output_path = Path(args.output)

    if not content_path.exists():
        print(f"ERROR: content file not found: {content_path}", file=sys.stderr)
        return 1
    content = json.loads(content_path.read_text(encoding="utf-8"))

    warnings = []
    if template_path.exists():
        doc = Document(str(template_path))
        warnings = apply_content(doc, content)
    else:
        if "fresh_build" not in content:
            print(f"ERROR: no template at {template_path} and no 'fresh_build' in content JSON.",
                  file=sys.stderr)
            return 1
        doc = build_fresh(content["fresh_build"])

    output_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(output_path))

    print(f"Wrote {output_path}")
    for w in warnings:
        print(f"  WARNING: {w}", file=sys.stderr)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
