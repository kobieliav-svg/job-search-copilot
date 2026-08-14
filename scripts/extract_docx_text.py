#!/usr/bin/env python3
"""Dump plain text from a .docx resume so it can be read during onboarding.

Usage:
    python scripts/extract_docx_text.py /path/to/resume.docx
"""
import sys

from docx import Document


def extract(path: str) -> str:
    doc = Document(path)
    parts = []
    for para in doc.paragraphs:
        text = para.text.strip()
        if text:
            parts.append(text)
    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                parts.append(" | ".join(cells))
    return "\n".join(parts)


def main() -> int:
    if len(sys.argv) != 2:
        print("Usage: extract_docx_text.py <path-to-docx>", file=sys.stderr)
        return 1
    print(extract(sys.argv[1]))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
