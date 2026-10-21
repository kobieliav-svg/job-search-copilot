#!/usr/bin/env python3
"""Mirror the dashboard's activity log into the jobs file, sheet "Activity".

The dashboard keeps the log (messages, referrals, calls, follow-ups); this keeps a backup copy in
pipeline/jobs.xlsx so nothing lives only in the browser. Rows are keyed by the activity id:
an existing id is updated in place, a new id is appended. Nothing is ever deleted.

    python scripts/activity_log.py sync --file activities.json
        activities.json: a JSON list of the dashboard's activity documents
        prints: {"ok": true, "written": n, "ids": [...]}
"""
import argparse
import json
import sys
from pathlib import Path

from openpyxl.styles import Font

import pipeline

SHEET = "Activity"
COLS = ["id", "date", "type", "person", "person_url", "company", "job_title", "job_url",
        "note", "follow_up_on", "status", "created_at", "notes", "outcome", "closed_at"]
# "notes" (the thread of replies and notes added after the action, as JSON),
# "outcome" and "closed_at" (how it ended) record how each action ended.


def sync(items):
    wb, _ = pipeline._load()
    if SHEET in wb.sheetnames:
        ws = wb[SHEET]
    else:
        ws = wb.create_sheet(SHEET)
        for i, c in enumerate(COLS, start=1):
            ws.cell(row=1, column=i, value=c).font = Font(bold=True)
        widths = {"A": 18, "B": 12, "C": 18, "D": 24, "E": 34, "F": 28, "G": 34, "H": 34, "I": 60, "J": 12, "K": 10, "L": 22}
        for k, v in widths.items():
            ws.column_dimensions[k].width = v
        ws.freeze_panes = "A2"
    for i, c in enumerate(COLS, start=1):  # older sheets lack the newer columns
        if ws.cell(row=1, column=i).value != c:
            ws.cell(row=1, column=i, value=c).font = Font(bold=True)
    where = {ws.cell(row=r, column=1).value: r for r in range(2, ws.max_row + 1) if ws.cell(row=r, column=1).value}
    ids = []
    for a in items:
        aid = str(a.get("id") or "").strip()
        if not aid:
            continue
        row = where.get(aid) or ws.max_row + 1
        where[aid] = row
        for i, c in enumerate(COLS, start=1):
            v = a.get(c, "")
            if isinstance(v, (list, dict)):
                v = json.dumps(v, ensure_ascii=False)
            ws.cell(row=row, column=i, value="" if v is None else str(v))
        ids.append(aid)
    pipeline._save(wb)
    return ids


def main():
    ap = argparse.ArgumentParser()
    sub = ap.add_subparsers(dest="cmd", required=True)
    s = sub.add_parser("sync"); s.add_argument("--file", required=True)
    a = ap.parse_args()
    items = json.loads(Path(a.file).read_text(encoding="utf-8"))
    try:
        ids = sync(items if isinstance(items, list) else [items])
    except (PermissionError, OSError) as e:
        print(json.dumps({"ok": False, "error": str(e)}))
        return 1
    print(json.dumps({"ok": True, "written": len(ids), "ids": ids}))
    return 0


if __name__ == "__main__":
    sys.exit(main())
