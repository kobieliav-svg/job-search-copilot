#!/usr/bin/env python3
"""Fill the user's real resume template with tailored content for one job posting.

Does NOT build a resume from scratch. Opens templates/resume_master.docx (the user's
actual resume, copied in during onboarding) and edits section text in place, so fonts,
margins, and layout stay exactly as the user already has them.

Content JSON shape (all keys optional -- only include what you're changing for this job):
{
  "paragraph_replacements": {
    "Summary": "One tailored summary line for this posting."
  },
  "bullet_replacements": {
    "Software Engineer — Acme Inc. (2021-2024)": [
      "Tailored bullet 1, pulled from FACTS.md, emphasizing what this job wants.",
      "Tailored bullet 2."
    ]
  },
  "placeholders": {
    "{{TARGET_ROLE}}": "Senior Data Analyst"
  }
}

Heading text in "paragraph_replacements" / "bullet_replacements" keys must match the
*exact* text of the heading paragraph already in the user's resume (case-insensitive).

Usage:
    python scripts/generate_resume.py --template templates/resume_master.docx \\
        --content output/tmp_acme.json --output "output/Acme_DataAnalyst_Resume.docx"
"""
import argparse
import json
import sys
from pathlib import Path

from docx import Document

import docx_editor


def apply_content(doc: Document, content: dict) -> list[str]:
    warnings = []

    for heading, text in content.get("paragraph_replacements", {}).items():
        if not docx_editor.replace_paragraph_text(doc, heading, text):
            warnings.append(f"Could not find heading '{heading}' for a paragraph replacement.")

    for heading, bullets in content.get("bullet_replacements", {}).items():
        if not docx_editor.replace_bullets(doc, heading, bullets):
            warnings.append(f"Could not find heading '{heading}' for a bullet replacement.")

    for placeholder, value in content.get("placeholders", {}).items():
        n = docx_editor.replace_placeholder(doc, placeholder, value)
        if n == 0:
            warnings.append(f"Placeholder '{placeholder}' not found anywhere in the document.")

    return warnings


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("--template", required=True)
    ap.add_argument("--content", required=True, help="Path to the content JSON file.")
    ap.add_argument("--output", required=True)
    args = ap.parse_args()

    template_path = Path(args.template)
    content_path = Path(args.content)
    output_path = Path(args.output)

    if not template_path.exists():
        print(f"ERROR: template not found: {template_path}", file=sys.stderr)
        return 1
    if not content_path.exists():
        print(f"ERROR: content file not found: {content_path}", file=sys.stderr)
        return 1

    content = json.loads(content_path.read_text(encoding="utf-8"))
    doc = Document(str(template_path))
    warnings = apply_content(doc, content)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(output_path))

    print(f"Wrote {output_path}")
    for w in warnings:
        print(f"  WARNING: {w}", file=sys.stderr)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
